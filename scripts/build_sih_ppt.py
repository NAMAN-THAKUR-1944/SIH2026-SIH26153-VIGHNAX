"""Fill the official SIH 2026 idea-presentation template for PS SIH26153.

    python scripts/build_sih_ppt.py [--template SIH2026-IDEA-Presentation-Format.pptx]
        ->  docs/VIGHNAX_SIH_Idea_Presentation.pptx   (export it to PDF for the SIH portal)

The template is https://sih.gov.in/letters/2026/SIH2026-IDEA-Presentation-Format.pptx (downloaded when no
path is given). As the SIH instructions require, its slide headings and pointers are kept, the deck stays
within six slides including the title slide, and the instructions slide is removed. Numbers come from
results/metrics.json via build_docs.py.
"""

from __future__ import annotations

import argparse
import copy
import os
import sys
import tempfile
import urllib.request

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_docs import BY, DATASET_NAMES, LEAD_L, LEAD_W, NON_TEMPORAL, ONSETS, ROOT, auc, f, mac  # noqa: E402

TEMPLATE_URL = "https://sih.gov.in/letters/2026/SIH2026-IDEA-Presentation-Format.pptx"
REPO_URL = "https://github.com/NAMAN-THAKUR-1944/SIH2026-SIH26153-VIGHNAX"
TEAM = {"ps_id": "SIH26153", "ps_title": "AI based Network Attack Forecasting from Network Traffic Data",
        "theme": "Blockchain & Cybersecurity", "category": "Software", "team_id": "127364", "team_name": "VIGHNAX"}
IDEA_TITLE = "VIGHNAX: a world model that forecasts network attacks before they complete"
BLUE = RGBColor(0x1F, 0x4E, 0x9A)
INK = RGBColor(0x20, 0x20, 0x20)
MUTED = RGBColor(0x55, 0x55, 0x55)
FONT = "Arial"


# ------------------------------------------------------------------ helpers
def shape(slide, name_prefix):
    return next(s for s in slide.shapes if s.name.startswith(name_prefix))


def set_text(sh, text):
    """Replace a shape's text, keeping the formatting of its first run."""
    p = sh.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in sh.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def bullet(p, char, level):
    pPr = p._p.get_or_add_pPr()
    indent = 228600  # 0.25 in
    pPr.set("marL", str(indent * (level + 1)))
    pPr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if char is None:
        none = pPr.makeelement(qn("a:buNone"), {})
        after = next((el for el in pPr if el.tag in (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))), None)
        pPr.append(none) if after is None else after.addprevious(none)
        return
    face, char = ("Wingdings", "v") if char == "❖" else ("Arial", char)   # the template draws ❖ as Wingdings "v"
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": face})
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": char})
    # DrawingML order: bullets come before tabLst / defRPr / extLst, or PowerPoint ignores them
    after = next((el for el in pPr if el.tag in (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))), None)
    for el in (bu_font, bu_char):
        if after is None:
            pPr.append(el)
        else:
            after.addprevious(el)


def para(tf, text, size, bold=False, color=INK, char="•", level=0, space_before=0, first=False, runs=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    bullet(p, char, level)
    if space_before:
        p.space_before = Pt(space_before)
    for t, b in (runs or [(text, bold)]):
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), b, FONT
        r.font.color.rgb = color
    return p


def body_box(slide, name_prefix, left, top, width, height):
    """Reuse the template's body text box (so its pointer styling stays), moved to the given area and cleared."""
    sh = shape(slide, name_prefix)
    sh.left, sh.top, sh.width, sh.height = Inches(left), Inches(top), Inches(width), Inches(height)
    tf = sh.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf


def section(tf, pointer, items, size=13, head=15, first=False):
    """One template pointer (kept verbatim, bold) followed by our points."""
    para(tf, pointer, head, bold=True, color=BLUE, char="❖", first=first, space_before=0 if first else 8)
    for it in items:
        if isinstance(it, tuple):
            para(tf, "", size, runs=[(it[0], True), (it[1], False)], level=1, space_before=2)
        else:
            para(tf, it, size, level=1, space_before=2)


def delete_slide(prs, index):
    sld_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld_id.get(qn("r:id")))
    prs.slides._sldIdLst.remove(sld_id)


# ------------------------------------------------------------------ slides
def title_slide(s):
    box = shape(s, "TextBox 9")
    tf = box.text_frame
    fields = [("Problem Statement ID – ", TEAM["ps_id"]), ("Problem Statement Title- ", TEAM["ps_title"]),
              ("Theme- ", TEAM["theme"]), ("PS Category- ", TEAM["category"]), ("Team ID- ", TEAM["team_id"]),
              ("Team Name (Registered on portal)- ", TEAM["team_name"])]
    template_p = copy.deepcopy(tf.paragraphs[1]._p)          # a bulleted field paragraph
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    txBody = tf._txBody
    for label, value in fields:
        p = copy.deepcopy(template_p)
        for r in p.findall(qn("a:r")):
            p.remove(r)
        txBody.append(p)
    for (label, value), p in zip(fields, tf.paragraphs):
        p.space_before = Pt(6)
        for t, b in ((label, True), (value, False)):
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(20), b, FONT
    box.width = Inches(6.9)


def idea_slide(s):
    title = s.shapes.title
    set_text(title, IDEA_TITLE)
    title.left, title.width, title.top, title.height = Inches(1.95), Inches(8.6), Inches(0.25), Inches(0.95)
    title.text_frame.word_wrap = True
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r in title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(22)
    tf = body_box(s, "TextBox 8", 0.45, 1.3, 12.4, 5.55)
    para(tf, "Proposed Solution (Describe your Idea/Solution/Prototype)", 19, bold=True, color=BLUE, char="❖", first=True)
    section(tf, "Detailed explanation of the proposed solution", [
        "A world model (recurrent state-space model) learns how a network's state evolves: every 10 s it reads a "
        "78-feature state built from flow records and/or packet captures",
        "It simulates 32 possible futures 60 s ahead: the probability that an attack is about to progress, with an "
        "uncertainty band, and the MITRE ATT&CK stage it is heading to",
        "Every alert is explained (Integrated Gradients) and traced to the responsible hosts and flows; "
        "working prototype: offline dashboard, PCAP / NetFlow-style CSV in, forecasts out, window by window",
    ], size=15, head=16.5)
    section(tf, "How it addresses the problem", [
        f"Moves from detection after the fact to early warning: {LEAD_W} of {ONSETS} attack onsets in unseen test "
        f"traffic were forecast before they began (logistic-regression baseline: {LEAD_L})",
        "One model trained on all seven datasets named in the PS and tested only on days, captures and botnet "
        "families it never saw",
    ], size=15, head=16.5)
    section(tf, "Innovation and uniqueness of the solution", [
        "Generative 'what-if' simulation of the network instead of a per-flow classifier",
        "Answers who, not only what: removing each host and re-running the forecast ranks the hosts behind an alert",
        "Air-gap ready: runs fully offline on a laptop CPU, SHA-256-verified model weights, one-click incident report",
    ], size=15, head=16.5)


def technical_slide(s):
    tf = body_box(s, "TextBox 8", 0.45, 1.25, 12.4, 1.95)
    section(tf, "Technologies to be used (e.g. programming languages, frameworks, hardware)", [
        ("Model: ", "Python, PyTorch (world model), scikit-learn (baseline); dpkt / Scapy for PCAP, pandas / NumPy"),
        ("Interface: ", "Flask + Chart.js dashboard with bundled fonts and scripts, no internet or cloud API"),
        ("Hardware: ", "any laptop CPU (~8 ms per 10 s window); a GPU only speeds up training"),
    ], size=12.5, head=14, first=True)
    para(tf, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", 14, bold=True,
         color=BLUE, char="❖", space_before=8)
    s.shapes.add_picture(os.path.join(ROOT, "docs", "images", "architecture.png"), Inches(1.1), Inches(2.75), width=Inches(11.1))


def feasibility_slide(s):
    tf = body_box(s, "TextBox 8", 0.45, 1.25, 7.55, 5.6)
    section(tf, "Analysis of the feasibility of the idea", [
        "Working prototype, trained weights and full benchmark are already built and public on GitHub",
        f"Unseen test data, macro over six temporal datasets: forecast AUROC {f(mac('forecast', 'world_model'))} vs "
        f"{f(mac('forecast', 'baseline_lr'))} (LR), detection {f(mac('detect', 'world_model'))} vs {f(mac('detect', 'baseline_lr'))}",
        "Throughput: one laptop CPU core processes traffic over 1,000x faster than real time",
    ], size=13, head=15, first=True)
    section(tf, "Potential challenges and risks", [
        "Weak on LANL (authentication-based red team) and behind LR on DARPA 1999; reported openly",
        "More false alarms on networks the model was not calibrated on; concept drift over time",
        "Encrypted traffic hides payloads",
    ], size=13, head=15)
    section(tf, "Strategies for overcoming these challenges", [
        "Per-network alert calibration on the site's own traffic (site profiles already built in)",
        "Authentication-log and per-host modelling for LANL-type attacks; periodic retraining",
        "Flow-level features need no payload, so the model works on encrypted traffic and on flow-only input",
    ], size=13, head=15)
    rows = [("Unseen test data", "Forecast", "Detection")]
    for d in BY:
        fc = "n/a" if d in NON_TEMPORAL else f"{f(auc(d, 'forecast', 'world_model'))} / {f(auc(d, 'forecast', 'baseline_lr'))}"
        rows.append((DATASET_NAMES[d], fc, f"{f(auc(d, 'detect', 'world_model'))} / {f(auc(d, 'detect', 'baseline_lr'))}"))
    rows.append(("Macro average", f"{f(mac('forecast', 'world_model'))} / {f(mac('forecast', 'baseline_lr'))}",
                 f"{f(mac('detect', 'world_model'))} / {f(mac('detect', 'baseline_lr'))}"))
    tbl = s.shapes.add_table(len(rows), 3, Inches(8.25), Inches(1.45), Inches(4.75), Inches(0.3 * len(rows))).table
    tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(2.05), Inches(1.35), Inches(1.35)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = val
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.runs[0]
            r.font.size, r.font.name = Pt(11), FONT
            r.font.bold = i == 0 or i == len(rows) - 1
            c.margin_top = c.margin_bottom = Inches(0.03)
    cap = s.shapes.add_textbox(Inches(8.25), Inches(1.5 + 0.3 * len(rows)), Inches(4.75), Inches(0.5)).text_frame
    cap.word_wrap = True
    para(cap, "AUROC, world model / logistic-regression baseline on the same features. CICIoT2023 episodes are "
              "composed, so only detection is scored.", 9.5, color=MUTED, char=None, first=True)


def impact_slide(s):
    tf = body_box(s, "TextBox 8", 0.45, 1.25, 6.9, 5.6)
    section(tf, "Potential impact on the target audience", [
        "NTRO, NCIIPC, CERT-In and critical-infrastructure SOCs get a warning before an attack completes, "
        "with time to isolate hosts",
        "Every alert arrives with its reasons and the responsible hosts and flows, so triage is faster",
        "Runs air-gapped on ordinary hardware, including classified and low-resource networks",
    ], size=13.5, head=15, first=True)
    section(tf, "Benefits of the solution (social, economic, environmental, etc.)", [
        ("Security: ", "earlier containment of botnets, DDoS and lateral movement on national infrastructure"),
        ("Social: ", "protects power, banking, health and e-governance services citizens rely on"),
        ("Economic: ", "open source, no licence or cloud cost; one CPU core covers a network segment"),
        ("Sovereignty: ", "indigenous, auditable AI; verified weights; no data leaves the network"),
    ], size=13.5, head=15)
    s.shapes.add_picture(os.path.join(ROOT, "docs", "images", "dashboard_light.png"), Inches(7.55), Inches(1.55), width=Inches(5.55))
    cap = s.shapes.add_textbox(Inches(7.55), Inches(4.6), Inches(5.55), Inches(0.6)).text_frame
    cap.word_wrap = True
    para(cap, "Working prototype: a real CIC-IDS2017 day never seen in training. Risk stays low for six minutes, "
              "then goes critical as the DDoS starts; the baseline (dashed) raises false spikes.", 10, color=MUTED, char=None, first=True)


def references_slide(s):
    tf = body_box(s, "TextBox 8", 0.45, 1.25, 12.4, 5.6)
    section(tf, "Details / Links of the reference and research work", [
        ("Source code, weights, benchmark: ", REPO_URL),
        ("World models: ", "Ha & Schmidhuber, World Models (2018); Hafner et al., Learning Latent Dynamics for Planning "
                           "from Pixels, ICML 2019 (recurrent state-space model)"),
        ("Explainability: ", "Sundararajan, Taly & Yan, Axiomatic Attribution for Deep Networks, ICML 2017 (Integrated Gradients)"),
        ("Attack stages: ", "MITRE ATT&CK Enterprise matrix, attack.mitre.org"),
        ("Datasets: ", "CTU-13 (Garcia et al., 2014); CIC-IDS2017 (Sharafaldin et al., 2018); CSE-CIC-IDS2018 "
                       "(CIC & AWS, 2018); UNSW-NB15 (Moustafa & Slay, 2015)"),
        ("", "LANL cyber1 (Kent, 2015); DARPA 1999 (Lippmann et al., 2000); CICIoT2023 (Neto et al., 2023)"),
        ("Guidance in the PS: ", "NCIIPC, nciipc.gov.in"),
        ("Reproducibility: ", "one configuration (configs/default.yaml, seed 42); python evaluate.py regenerates every "
                              "number in this deck from the public datasets"),
    ], size=15.5, head=17, first=True)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", help="path to SIH2026-IDEA-Presentation-Format.pptx (downloaded if omitted)")
    ap.add_argument("--out", default=os.path.join(ROOT, "docs", "VIGHNAX_SIH_Idea_Presentation.pptx"))
    a = ap.parse_args()
    tpl = a.template
    if not tpl:
        tpl = os.path.join(tempfile.gettempdir(), "SIH2026-IDEA-Presentation-Format.pptx")
        urllib.request.urlretrieve(TEMPLATE_URL, tpl)
    prs = Presentation(tpl)
    slides = list(prs.slides)
    title_slide(slides[0])
    for sl in slides[1:6]:
        oval = next(sh for sh in sl.shapes if sh.name.startswith("Oval"))
        set_text(oval, TEAM["team_name"])
        for r in oval.text_frame.paragraphs[0].runs:
            r.font.size, r.font.bold = Pt(12), True
    idea_slide(slides[1])
    technical_slide(slides[2])
    feasibility_slide(slides[3])
    impact_slide(slides[4])
    references_slide(slides[5])
    delete_slide(prs, 6)          # "Important instructions" slide (the template says to delete it before upload)
    prs.save(a.out)
    print("wrote", os.path.relpath(a.out, ROOT))


if __name__ == "__main__":
    main()
